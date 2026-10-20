#!/usr/bin/env python3
"""
Price Intelligence / Renegotiation Agent

Pipeline:
  STEP 1   Load ingredients from file (.txt / .md / .json)
  STEP 2   Web search (Tavily preferred, Google News RSS fallback)
  STEP 3   Select, scrape, and clean source pages
  STEP 3b  Official benchmark lookup (World Bank Pink Sheet, best effort)
  STEP 4   LLM analyst (Ollama locally, or OpenRouter)
  STEP 5   Renegotiation report (Markdown + optional PDF) + audit JSON

Usage:
  python 04_Ingredients_AI.py
  python 04_Ingredients_AI.py --backend openrouter --model openai/gpt-4o-mini -v
  python 04_Ingredients_AI.py --ingredients ../data/ingredients.txt --workers 4

Environment:
  TAVILY_API_KEY       Tavily search key (strongly recommended)
  OPENROUTER_API_KEY   OpenRouter backend key
Optional env overrides:
  OPENROUTER_MODEL, OLLAMA_MODEL, LLM_BACKEND, MAX_WORKERS

Optional packages (each feature degrades gracefully when missing):
  pip install requests beautifulsoup4        # core
  pip install tavily-python                  # Tavily search
  pip install ollama                         # local LLM backend
  pip install openai                         # OpenRouter backend
  pip install markdown weasyprint            # PDF export (weasyprint needs GTK runtime)
  pip install openpyxl                       # World Bank Pink Sheet benchmark
"""

import argparse
import json
import logging
import os
import re
import time
import xml.etree.ElementTree as ET
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from email.utils import parsedate_to_datetime
from pathlib import Path
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

# ----------------------------------------------------------------------------
# CONFIGURATION (defaults; CLI flags / env vars override)
# ----------------------------------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = (BASE_DIR / ".." / "data").resolve()
REPORTS_DIR = (BASE_DIR / ".." / "reports").resolve()
DEFAULT_INGREDIENTS_FILE = DATA_DIR / "ingredients.txt"
CACHE_DIR = DATA_DIR / "cache"

SOURCE_WINDOW_DAYS = 180       # ignore sources older than this
PINK_SHEET_TTL_DAYS = 7        # re-download benchmark workbook at most weekly
JSON_RETRIES = 2               # attempts to get valid JSON out of the LLM
MAX_SELECT_SOURCES = 4         # sources scraped + fed to the LLM per ingredient

LLM_BACKENDS = ("ollama", "openrouter", "transformers_local")
OPENROUTER_BASE_URL = "https://openrouter.ai/api/v1"
DEFAULT_MODELS = {
    "ollama": os.getenv("OLLAMA_MODEL", "qwen2.5:7b-instruct"),
    "openrouter": os.getenv("OPENROUTER_MODEL", "openai/gpt-4o-mini"),
    "transformers_local": "microsoft/Phi-3-mini-4k-instruct",
}

VALID_TRENDS = {"UP", "DOWN", "STABLE", "UNKNOWN"}
VALID_CONFIDENCE = {"HIGH", "MEDIUM", "LOW"}
REQUIRED_KEYS = ("trend", "confidence", "specific_quotes", "summary")

log = logging.getLogger("price-agent")


def setup_logging(verbose: bool) -> None:
    logging.basicConfig(
        level=logging.DEBUG if verbose else logging.INFO,
        format="%(asctime)s %(levelname)-7s %(message)s",
        datefmt="%H:%M:%S",
    )


# Shared session: browser-like UA + automatic retries on transient failures
_HTTP = requests.Session()
_HTTP.headers.update({
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    )
})
_RETRIES = Retry(total=2, connect=2, read=2, status=2, backoff_factor=0.5,
                 status_forcelist=(429, 500, 502, 503, 504))
_ADAPTER = HTTPAdapter(max_retries=_RETRIES)
_HTTP.mount("http://", _ADAPTER)
_HTTP.mount("https://", _ADAPTER)

# ----------------------------------------------------------------------------
# STEP 1: Load Ingredients from File
# ----------------------------------------------------------------------------


def load_ingredients(filepath: Path) -> list[str]:
    """Load ingredient names from .txt / .json / .md."""
    ext = filepath.suffix.lower()
    content = filepath.read_text(encoding="utf-8").strip()

    if ext == ".json":
        data = json.loads(content)
        if isinstance(data, dict) and "ingredients" in data:
            data = data["ingredients"]
        if isinstance(data, list):
            return [str(item).strip() for item in data if str(item).strip()]

    elif ext == ".md":
        # Markdown: bullet list or numbered list
        items = []
        for line in content.splitlines():
            line = line.strip()
            if line.startswith(("- ", "* ")):
                items.append(line[2:].strip())
            elif re.match(r"^\d+[.)]\s*", line):
                items.append(re.sub(r"^\d+[.)]\s*", "", line).strip())
        if items:
            return items

    # Fallback: comma-separated single line, else one ingredient per line
    if "," in content and "\n" not in content:
        return [part.strip() for part in content.split(",") if part.strip()]
    return [line.strip() for line in content.splitlines() if line.strip()]

# ----------------------------------------------------------------------------
# STEP 2: Web Search (Tavily preferred, fallback with Google News RSS)
# ----------------------------------------------------------------------------


def build_query(ingredient: str) -> str:
    """Recent price-trend query; the date window is computed, never hardcoded."""
    since = (datetime.now(timezone.utc) - timedelta(days=SOURCE_WINDOW_DAYS)).strftime("%Y-%m")
    return f"{ingredient} wholesale price trend since {since}"


def search_tavily(query: str) -> list[dict]:
    from tavily import TavilyClient

    client = TavilyClient(api_key=os.environ["TAVILY_API_KEY"])
    response = client.search(query, max_results=6, search_depth="advanced", time_range="year")
    results = []
    for hit in response.get("results", []):
        result = {
            "url": hit.get("url", ""),
            "title": hit.get("title", ""),
            "snippet": hit.get("content", ""),
            "date": hit.get("published_date"),
        }
        if result["url"]:
            results.append(result)
    return results


GOOGLE_NEWS_RSS = "https://news.google.com/rss/search"


def search_google_news(query: str) -> list[dict]:
    """Free fallback search via Google News RSS (no API key needed).

    Uses requests' params for proper URL encoding and stdlib XML parsing
    (no lxml dependency). Each item keeps its pubDate for recency filtering.
    """
    resp = _HTTP.get(
        GOOGLE_NEWS_RSS,
        params={"q": query, "hl": "en-US", "gl": "US", "ceid": "US:en"},
        timeout=15,
    )
    resp.raise_for_status()
    root = ET.fromstring(resp.content)
    results = []
    for item in root.iter("item"):
        raw_desc = item.findtext("description") or ""
        description = BeautifulSoup(raw_desc, "html.parser").get_text(" ", strip=True)
        results.append({
            "url": (item.findtext("link") or "").strip(),
            "title": (item.findtext("title") or "").strip(),
            "snippet": description,
            "date": item.findtext("pubDate"),
        })
        if len(results) >= 6:
            break
    return [r for r in results if r["url"]]


def search_ingredient_price(ingredient: str) -> list[dict]:
    """Return list of {url, title, snippet, date}; Tavily preferred, RSS fallback."""
    query = build_query(ingredient)

    if os.getenv("TAVILY_API_KEY"):
        try:
            results = search_tavily(query)
            if results:
                return results
            log.warning("Tavily returned no results for %r; falling back to Google News", ingredient)
        except Exception as exc:
            log.warning("Tavily failed for %r (%s); falling back to Google News", ingredient, exc)
    else:
        log.info("TAVILY_API_KEY not set - using Google News RSS fallback for %r", ingredient)

    try:
        return search_google_news(query)
    except Exception as exc:
        log.error("Google News RSS fallback failed for %r: %s", ingredient, exc)
        return []


# ---- Date parsing / recency filtering --------------------------------------


def parse_date(raw) -> datetime | None:
    """Parse RFC 822 (RSS), ISO 8601, or 'Mar 14, 2026' style dates."""
    if not raw:
        return None
    text = str(raw).strip()
    try:
        return parsedate_to_datetime(text)
    except (TypeError, ValueError):
        pass
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        pass
    for fmt in ("%b %d, %Y", "%d %b %Y", "%B %d, %Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def is_recent(published) -> bool:
    """True for recent items AND items whose date is missing/unparseable."""
    dt = parse_date(published)
    if dt is None:
        return True
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    cutoff = datetime.now(timezone.utc) - timedelta(days=SOURCE_WINDOW_DAYS)
    return dt >= cutoff


# ---- Source quality: trust scoring, de-duplication, selection ---------------

DOMAIN_TRUST = {
    "worldbank.org": 5, "imf.org": 5, "fao.org": 4, "amis.org": 4, "oecd.org": 4,
    "usda.gov": 5, "ers.usda.gov": 5, "bls.gov": 5, "usitc.gov": 4,
    "tradingeconomics.com": 4, "indexmundi.com": 4, "statista.com": 3,
    "reuters.com": 4, "bloomberg.com": 4, "ft.com": 4, "wsj.com": 3,
    "agriculture.com": 3, "farmdocdaily.illinois.edu": 4, "commodityonline.com": 3,
    "rabobank.com": 3, "fitchratings.com": 3,
}
_DEFAULT_TRUST = 2


def _registrable_domain(url_or_domain: str) -> str:
    """Coarse registrable domain: last two labels (aware of a few multi-part TLDs)."""
    if "://" not in url_or_domain:
        url_or_domain = f"https://{url_or_domain}"
    host = urlparse(url_or_domain).netloc.split("@")[-1].split(":")[0].lower()
    labels = [part for part in host.split(".") if part]
    if len(labels) >= 3 and ".".join(labels[-2:]) in {"co.uk", "com.br", "co.jp", "com.au", "gov.uk"}:
        return ".".join(labels[-3:])
    return ".".join(labels[-2:]) if len(labels) >= 2 else host


def trust_score(url: str) -> int:
    return DOMAIN_TRUST.get(_registrable_domain(url), _DEFAULT_TRUST)


def select_sources(results: list[dict], k: int = MAX_SELECT_SOURCES,
                   max_per_domain: int = 2) -> list[dict]:
    """De-duplicate, prefer trusted + recent domains, cap per-domain spread."""
    dated = [r for r in results if r.get("date") and is_recent(r["date"])]
    pool = dated if len(dated) >= k else results
    ranked = sorted(
        pool,
        key=lambda r: (-trust_score(r["url"]), -int(is_recent(r.get("date")))),
    )
    picked: list[dict] = []
    seen_urls: set[str] = set()
    per_domain: dict[str, int] = {}
    for result in ranked:
        if result["url"] in seen_urls:
            continue
        domain = _registrable_domain(result["url"])
        if per_domain.get(domain, 0) >= max_per_domain:
            continue
        seen_urls.add(result["url"])
        per_domain[domain] = per_domain.get(domain, 0) + 1
        picked.append(result)
        if len(picked) == k:
            break
    return picked

# ----------------------------------------------------------------------------
# STEP 3: Scrape and Clean Webpage Text
# ----------------------------------------------------------------------------


def fetch_page_text(url: str, max_chars: int = 4000) -> str | None:
    """Fetch and extract main text from a URL. None on failure (never an error string)."""
    try:
        resp = _HTTP.get(url, timeout=20)
        resp.raise_for_status()
        soup = BeautifulSoup(resp.content, "html.parser")
        for tag in soup(["script", "style", "nav", "footer", "header", "aside", "form", "noscript"]):
            tag.decompose()
        text = re.sub(r"\s+", " ", soup.get_text(" ", strip=True))
        return text[:max_chars] or None
    except Exception as exc:
        log.debug("Failed to fetch %s: %s", url, exc)
        return None


def _date_label(published) -> str:
    dt = parse_date(published)
    if dt:
        return dt.strftime("%Y-%m-%d")
    return str(published) if published else "unknown date"

# ----------------------------------------------------------------------------
# STEP 3b: Official benchmark data (World Bank "Pink Sheet", best effort)
#
# Structured prices beat news prose. If openpyxl is installed we download the
# World Bank monthly commodity workbook (cached weekly) and look up the
# ingredient via COMMODITY_ALIASES. Any failure just means "no benchmark":
# the news-based analysis carries on regardless.
# ----------------------------------------------------------------------------

WB_LANDING_URL = "https://www.worldbank.org/en/research/commodity-markets"
PINK_SHEET_FILE = "wb_pink_sheet_monthly.xlsx"

# ingredient word(s) -> commodity keyword to find in the Pink Sheet.
# Matched with word boundaries, longest alias wins.
COMMODITY_ALIASES = {
    "coffee": "Coffee", "cocoa": "Cocoa", "tea": "Tea", "sugar": "Sugar",
    "wheat flour": "Wheat", "flour": "Wheat", "wheat": "Wheat",
    "maize": "Maize", "corn": "Maize", "rice": "Rice",
    "soybean": "Soybeans", "soy beans": "Soybeans", "soy": "Soybeans",
    "olive oil": None,                # not in the Pink Sheet -> no benchmark
    "palm oil": "Palm oil", "vegetable oil": "Palm oil",      # proxy
    "canola": "Rapeseed oil", "rapeseed": "Rapeseed oil",
    "sunflower oil": "Sunflower oil", "sunflower": "Sunflower oil",
    "peanut": "Groundnuts", "groundnut": "Groundnuts",
    "beef": "Beef", "chicken": "Chicken", "pork": "Pork", "lamb": "Lamb",
    "mutton": "Sheep meat", "butter": "Butter", "cheese": "Cheese",
    "milk powder": "Powdered milk", "whey": "Whey", "milk": "Milk",
    "cotton": "Cotton", "rubber": "Rubber",
    "aluminum": "Aluminium", "aluminium": "Aluminium", "copper": "Copper",
    "nickel": "Nickel", "zinc": "Zinc", "tin": "Tin", "iron ore": "Iron ore",
    "tuna": "Tuna", "salmon": "Salmon", "shrimp": "Shrimp",
    "banana": "Banana", "orange juice": "Orange juice", "orange": "Orange",
    "eggs": None, "egg": None, "salt": None,  # not covered by the Pink Sheet
}

_PINK_SHEET_TABLE: dict | None = None
_PINK_SHEET_LOADED = False

_MONTH_LABEL_RE = re.compile(
    r"^(?:[A-Za-z]{3}[-/ ]?\d{2,4}|\d{4}[-M]\d{1,2})$"   # Feb-26 / Feb 2026 / 2026-02
)


def download_pink_sheet(max_age_days: int = PINK_SHEET_TTL_DAYS) -> Path | None:
    """Download (weekly-cached) the World Bank monthly commodity XLSX."""
    cache = DATA_DIR / PINK_SHEET_FILE
    try:
        if cache.exists() and time.time() - cache.stat().st_mtime < max_age_days * 86400:
            return cache
        landing = _HTTP.get(WB_LANDING_URL, timeout=20)
        link = next(
            (urljoin(landing.url, anchor["href"])
             for anchor in BeautifulSoup(landing.content, "html.parser").find_all("a", href=True)
             if "Monthly.xlsx" in anchor["href"]),
            None,
        )
        if not link:
            log.warning("World Bank Pink Sheet link not found on landing page; skipping benchmark")
            return None
        payload = _HTTP.get(link, timeout=90)
        payload.raise_for_status()
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        cache.write_bytes(payload.content)
        log.info("Downloaded World Bank Pink Sheet -> %s", cache)
        return cache
    except Exception as exc:
        log.warning("World Bank Pink Sheet download failed (%s); skipping benchmark", exc)
        return None


def _load_pink_sheet() -> dict | None:
    """Parse the workbook once per run. Returns {row_name: (unit, values, labels)}."""
    global _PINK_SHEET_TABLE, _PINK_SHEET_LOADED
    if _PINK_SHEET_LOADED:
        return _PINK_SHEET_TABLE
    _PINK_SHEET_LOADED = True

    path = download_pink_sheet()
    if path is None:
        return None
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = next(
            (ws for ws in workbook.worksheets if "monthly" in ws.title.lower()),
            workbook.active,
        )
        month_labels: dict[int, str] = {}
        table: dict[str, tuple[str, list[float], list[str]]] = {}
        for row in sheet.iter_rows(values_only=True):
            cells = list(row)
            # Detect the month-label header row: label first appears in col >= 2
            if len(cells) > 2 and len([c for c in cells[2:] if isinstance(c, str)
                                       and _MONTH_LABEL_RE.match(c.strip())]) >= 6:
                month_labels = {
                    i: c.strip() for i, c in enumerate(cells[2:], start=2)
                    if isinstance(c, str) and _MONTH_LABEL_RE.match(c.strip())
                }
                continue
            # Data row: text name in col 0
            if not (cells and isinstance(cells[0], str) and cells[0].strip()):
                continue
            name = cells[0].strip()
            unit = cells[1].strip() if len(cells) > 1 and isinstance(cells[1], str) else ""
            values, labels = [], []
            for i, value in enumerate(cells[2:], start=2):
                if month_labels and i > max(month_labels):
                    break   # stop at summary columns (kept out by Col-i limit)
                if not month_labels and not isinstance(value, (int, float)):
                    break   # no header row found: stop at first non-numeric cell
                if isinstance(value, (int, float)):
                    values.append(float(value))
                    labels.append(month_labels.get(i, ""))
                elif isinstance(value, str):
                    break
            if len(values) >= 6:
                table[name] = (unit, values, labels)
        workbook.close()
        _PINK_SHEET_TABLE = table or None
        log.info("Pink Sheet parsed: %d commodity rows", len(table))
    except Exception as exc:
        log.warning("Pink Sheet parsing failed (%s); skipping benchmark", exc)
        _PINK_SHEET_TABLE = None
    return _PINK_SHEET_TABLE


def wb_benchmark(ingredient: str) -> dict | None:
    """Best-effort official price benchmark for an ingredient."""
    table = _load_pink_sheet()
    if not table:
        return None

    ingredient_lower = ingredient.lower()
    keyword = ingredient_lower
    for alias, commodity in sorted(COMMODITY_ALIASES.items(), key=lambda kv: -len(kv[0])):
        if re.search(rf"\b{re.escape(alias)}\b", ingredient_lower):
            keyword = (commodity or "").lower()
            break
    if not keyword:
        return None
    matches = [item for item in table.items() if keyword in item[0].lower()]
    if not matches:
        return None
    # Prefer 'world' reference quotes over regional ones, then sheet order
    matches.sort(key=lambda item: ("world" not in item[0].lower(),))
    name, (unit, raw_values, raw_labels) = matches[0]

    pairs = [(lbl, val) for lbl, val in zip(raw_labels, raw_values) if val == val and val > 0]
    if len(pairs) < 4:
        return None
    latest_label, latest = pairs[-1]
    change_3m = 100.0 * (latest / pairs[-4][1] - 1)
    change_6m = 100.0 * (latest / pairs[-7][1] - 1) if len(pairs) >= 7 else None
    return {
        "commodity": name,
        "unit": unit,
        "latest": round(latest, 3),
        "change_3m_pct": round(change_3m, 1),
        "change_6m_pct": round(change_6m, 1) if change_6m is not None else None,
        "as_of": latest_label or "latest available",
        "source": "World Bank Commodity Markets (Pink Sheet)",
    }

# ----------------------------------------------------------------------------
# STEP 4: LLM Analyst (Ollama locally, or OpenRouter)
# ----------------------------------------------------------------------------

ANALYST_SYSTEM_PROMPT = (
    "You are a meticulous commodity-pricing analyst. You only report prices and trends "
    "that are supported by the supplied context, and you answer strictly in the requested "
    "JSON format."
)

_LOCAL_PIPELINE = None  # lazy-loaded transformers pipeline


def call_llm(prompt: str, backend: str, model: str, json_mode: bool = True) -> str:
    messages = [
        {"role": "system", "content": ANALYST_SYSTEM_PROMPT},
        {"role": "user", "content": prompt},
    ]

    if backend == "ollama":
        import ollama

        kwargs = {"model": model, "messages": messages, "options": {"temperature": 0}}
        if json_mode:
            kwargs["format"] = "json"
        response = ollama.chat(**kwargs)
        return response["message"]["content"].strip()

    if backend == "openrouter":
        from openai import OpenAI

        client = OpenAI(base_url=OPENROUTER_BASE_URL, api_key=os.environ["OPENROUTER_API_KEY"])
        if json_mode:
            try:
                completion = client.chat.completions.create(
                    model=model, messages=messages, temperature=0,
                    response_format={"type": "json_object"},
                )
                return completion.choices[0].message.content.strip()
            except Exception as exc:
                # Some OpenRouter models reject response_format; degrade gracefully.
                log.debug("OpenRouter %s rejected json_object; retrying plain: %s", model, exc)
        completion = client.chat.completions.create(model=model, messages=messages, temperature=0)
        return completion.choices[0].message.content.strip()

    if backend == "transformers_local":
        global _LOCAL_PIPELINE
        from transformers import pipeline

        if _LOCAL_PIPELINE is None:
            try:
                _LOCAL_PIPELINE = pipeline("text-generation", model=model, device_map="auto")
            except Exception:
                _LOCAL_PIPELINE = pipeline("text-generation", model=model, device=-1)
        output = _LOCAL_PIPELINE(messages, max_new_tokens=600)
        return output[0]["generated_text"][-1]["content"].strip()

    raise ValueError(f"Unknown LLM backend: {backend}")


def _validate_analysis(data) -> str | None:
    """Return an error description, or None if the analysis dict is well-formed."""
    if not isinstance(data, dict):
        return "not a JSON object"
    missing = [key for key in REQUIRED_KEYS if key not in data]
    if missing:
        return f"missing keys: {', '.join(missing)}"
    if str(data["trend"]).strip().upper() not in VALID_TRENDS:
        return f"trend must be one of {sorted(VALID_TRENDS)}"
    if str(data["confidence"]).strip().upper() not in VALID_CONFIDENCE:
        return f"confidence must be one of {sorted(VALID_CONFIDENCE)}"
    if not isinstance(data["specific_quotes"], list):
        return "specific_quotes must be a list"
    return None


def _extract_json(raw: str) -> dict | None:
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        pass
    match = re.search(r"\{.*\}", raw, re.DOTALL)
    if match:
        try:
            return json.loads(match.group())
        except json.JSONDecodeError:
            return None
    return None


def analyze_price_trend(ingredient: str, results: list[dict], benchmark: dict | None,
                        backend: str, model: str) -> dict:
    """Use the LLM to extract price direction, quotes, and evidence."""
    selected = select_sources(results, k=MAX_SELECT_SOURCES)

    # Scrape in parallel; fall back to the search snippet when a page blocks us
    with ThreadPoolExecutor(max_workers=min(3, max(1, len(selected)))) as pool:
        pages = list(pool.map(lambda res: fetch_page_text(res["url"]), selected))

    context_blocks: list[str] = []
    sources_meta: list[dict] = []
    for result, page in zip(selected, pages):
        content = page or (result.get("snippet") or "").strip()
        if not content:
            continue   # failed fetch AND no snippet -> don't pollute the LLM context
        note = "" if page else " (page fetch failed; using search snippet)"
        date_label = _date_label(result.get("date"))
        context_blocks.append(
            f"--- SOURCE {len(context_blocks) + 1}{note}\n"
            f"Title: {result['title']}\n"
            f"URL: {result['url']}\n"
            f"Published: {date_label}\n"
            f"Content: {content[:1500]}"
        )
        sources_meta.append({"url": result["url"], "title": result["title"], "date": date_label})

    benchmark_note = ""
    if benchmark:
        six_month = (f", 6 months: {benchmark['change_6m_pct']:+.1f}%"
                     if benchmark.get("change_6m_pct") is not None else "")
        benchmark_note = (
            f"OFFICIAL BENCHMARK (authoritative; prefer this over news prose):\n"
            f"{benchmark['commodity']} ({benchmark['unit']}): latest {benchmark['latest']} "
            f"(as of {benchmark['as_of']}), 3-month change: {benchmark['change_3m_pct']:+.1f}%{six_month}"
        )

    if not context_blocks and not benchmark_note:
        return {
            "ingredient": ingredient, "trend": "UNKNOWN", "confidence": "LOW",
            "specific_quotes": [], "summary": "No usable search results or page content.",
            "sources": [], "benchmark": None,
        }

    prompt = f"""You are a commodity pricing analyst preparing intelligence for a procurement team
before contract renegotiation. Below are web excerpts for "{ingredient}"
{('plus an official benchmark. ' + chr(10) + benchmark_note) if benchmark_note else ''}

Using ONLY the information in the context above (never invent or guess prices), determine:
1. The overall price trend (UP, DOWN, STABLE, or UNKNOWN) over the last 3-6 months.
2. Specific price quotes found in the context, each with currency, unit, and date/date period.
3. A 2-3 sentence executive summary.

Rules:
- If the context contains no usable price information, return trend UNKNOWN, confidence LOW,
  an empty quote list, and a one-sentence explanation.
- Cite sources as SOURCE 1, SOURCE 2, ... in quotes; official benchmark data may be cited
  as "World Bank Pink Sheet".
- Quote format example: "215 USD per ton (Apr 2026, SOURCE 2)".

Return strictly this JSON (no markdown fences, no commentary):
{{
  "trend": "UP|DOWN|STABLE|UNKNOWN",
  "confidence": "HIGH|MEDIUM|LOW",
  "specific_quotes": ["quote with currency, unit and date", "..."],
  "summary": "your concise summary"
}}

Context ({len(context_blocks)} excerpt(s)):
{chr(10) + chr(10).join(context_blocks) if context_blocks else '(no web excerpts; official benchmark only)'}
"""

    raw = ""
    data = None
    last_error = None
    for attempt in range(1, JSON_RETRIES + 1):
        attempt_prompt = prompt
        if last_error:
            attempt_prompt += (f"\nYour previous reply was invalid ({last_error}). "
                               "Respond with ONLY the valid JSON object.")
        raw = call_llm(attempt_prompt, backend, model)
        candidate = _extract_json(raw)
        last_error = _validate_analysis(candidate)
        if last_error is None:
            data = candidate
            break
        log.debug("LLM JSON invalid for %r (attempt %d): %s", ingredient, attempt, last_error)

    if data is None:
        # Last resort: word-boundary heuristic so one bad parse doesn't kill the run
        log.warning("Falling back to heuristic trend extraction for %r", ingredient)
        match = re.search(r"\b(UP|DOWN|STABLE)\b", raw, re.IGNORECASE)
        data = {
            "trend": match.group(1).upper() if match else "UNKNOWN",
            "confidence": "LOW",
            "specific_quotes": [],
            "summary": raw.strip()[:300],
        }

    data["trend"] = str(data["trend"]).strip().upper()
    data["confidence"] = str(data["confidence"]).strip().upper()
    data["ingredient"] = ingredient
    data["sources"] = sources_meta
    data["benchmark"] = benchmark
    return data

# ----------------------------------------------------------------------------
# STEP 5: Generate Renegotiation Report (Markdown + optional PDF)
# ----------------------------------------------------------------------------

LEVERAGE_GUIDANCE = {
    "DOWN": ("Strong leverage",
             "Prices are falling - demand a 5-10% reduction from current contracts, "
             "citing recent market drops."),
    "UP": ("Limited leverage",
           "Prices are rising - consider locking in long-term fixed pricing or "
           "exploring alternative suppliers."),
    "STABLE": ("Neutral",
               "Prices are stable - renegotiate on volume, service levels, and "
               "relationship rather than market pressure."),
    "UNKNOWN": ("Unknown",
                "No reliable market signal found - verify prices manually before negotiating."),
}


def generate_report(analyses: list[dict], meta: dict) -> str:
    counts: dict[str, int] = {}
    for analysis in analyses:
        trend = analysis.get("trend", "UNKNOWN")
        counts[trend] = counts.get(trend, 0) + 1

    lines = [
        "# Procurement Price Renegotiation Report",
        "",
        f"**Generated:** {datetime.now().strftime('%Y-%m-%d %H:%M')}  ",
        f"**Analyst:** {meta['backend']} ({meta['model']}) | "
        f"Search window: last {SOURCE_WINDOW_DAYS} days",
        "",
        "## Executive Summary",
        "",
        f"Market signal across {len(analyses)} ingredient(s): "
        + ", ".join(f"{value} {trend}" for trend, value in sorted(counts.items())) + ".",
        "",
        "| Ingredient | Trend | Confidence | Leverage |",
        "|---|---|---|---|",
    ]
    for analysis in analyses:
        short, _ = LEVERAGE_GUIDANCE.get(analysis.get("trend", "UNKNOWN"),
                                         LEVERAGE_GUIDANCE["UNKNOWN"])
        lines.append(
            f"| {analysis['ingredient']} | {analysis.get('trend', 'UNKNOWN')} "
            f"| {analysis.get('confidence', 'LOW')} | {short} |"
        )

    lines += ["", "---", ""]
    for analysis in analyses:
        short, detail = LEVERAGE_GUIDANCE.get(analysis.get("trend", "UNKNOWN"),
                                              LEVERAGE_GUIDANCE["UNKNOWN"])
        lines += [f"## {analysis['ingredient'].upper()}", ""]

        bench = analysis.get("benchmark")
        if bench:
            six_month = (f" | 6 mo: {bench['change_6m_pct']:+.1f}%"
                         if bench.get("change_6m_pct") is not None else "")
            lines += [
                f"**Official benchmark:** {bench['commodity']} - latest "
                f"**{bench['latest']} {bench['unit']}** ({bench['as_of']}) "
                f"| 3 mo: {bench['change_3m_pct']:+.1f}%{six_month} "
                f"*(World Bank Pink Sheet)*",
                "",
            ]

        quotes = [q for q in analysis.get("specific_quotes", []) if q]
        lines += [
            f"**Trend:** {analysis.get('trend', 'UNKNOWN')} | "
            f"**Confidence:** {analysis.get('confidence', 'LOW')}",
            "",
            f"**Price quotes:** " + ("; ".join(quotes) if quotes else "_none found in sources_"),
            "",
            f"**Analysis:** {analysis.get('summary', 'N/A')}",
            "",
            f"**Leverage:** **{short}:** {detail}",
        ]

        sources = analysis.get("sources") or []
        if sources:
            lines += ["", "**Sources:**"]
            lines += [f"- [{source['title']}]({source['url']}) - {source['date']}"
                      for source in sources]

        lines += ["", "---", ""]

    lines.append("*Generated automatically by the Price Intelligence Agent. "
                 "Verify prices against official sources before negotiating.*")
    return "\n".join(lines)


# ---- Result cache (weekly) + audit trail ------------------------------------


def cache_path(ingredient: str) -> Path:
    iso = datetime.now(timezone.utc).isocalendar()
    slug = re.sub(r"[^a-z0-9]+", "-", ingredient.lower()).strip("-") or "ingredient"
    return CACHE_DIR / f"{iso[0]}-W{iso[1]:02d}_{slug}.json"


def process_ingredient(ingredient: str, args) -> dict:
    """Search + benchmark + analyze one ingredient, with a weekly result cache."""
    cache = cache_path(ingredient)
    if not args.no_cache:
        try:
            cached = json.loads(cache.read_text(encoding="utf-8"))
            if isinstance(cached, dict) and cached.get("trend"):
                log.info("[cache] Using cached analysis for %r "
                         "(weekly cache; --no-cache to refresh)", ingredient)
                return cached
        except (OSError, ValueError):
            pass

    log.info("Processing %r ...", ingredient)
    results = search_ingredient_price(ingredient)
    selected = select_sources(results, k=MAX_SELECT_SOURCES)
    log.info("  %d search result(s), %d selected for scraping", len(results), len(selected))

    benchmark = wb_benchmark(ingredient)
    if benchmark:
        log.info("  World Bank benchmark: %s = %s %s (%s)",
                 benchmark["commodity"], benchmark["latest"], benchmark["unit"],
                 benchmark["as_of"])
    else:
        log.info("  No official benchmark match; relying on web sources")

    analysis = analyze_price_trend(ingredient, results, benchmark, args.backend, args.model)
    log.info("  Trend: %s (confidence %s)", analysis.get("trend"), analysis.get("confidence"))

    if not args.no_cache:
        try:
            CACHE_DIR.mkdir(parents=True, exist_ok=True)
            cache.write_text(json.dumps(analysis, ensure_ascii=False, indent=2),
                             encoding="utf-8")
        except OSError as exc:
            log.debug("Could not write cache for %r: %s", ingredient, exc)
    return analysis


# ---- Optional PDF export -----------------------------------------------------


def export_pdf(md_path: Path) -> None:
    """Render the Markdown report to a styled PDF next to it."""
    try:
        import markdown
    except ImportError:
        log.info("PDF skipped (markdown not installed - pip install markdown weasyprint)")
        return
    try:
        from weasyprint import HTML
    except Exception as exc:
        log.info("PDF skipped (weasyprint unavailable; it needs the GTK runtime on "
                 "Windows): %s", exc)
        return

    body = markdown.markdown(md_path.read_text(encoding="utf-8"), extensions=["tables"])
    html = (
        "<html><head><meta charset='utf-8'><style>"
        "body{font-family:'Segoe UI',Arial,sans-serif;font-size:11pt;color:#222;margin:36px;}"
        "h1{font-size:20pt;color:#14365d;border-bottom:3px solid #14365d;padding-bottom:8px;}"
        "h2{font-size:15pt;color:#14365d;margin-top:28px;}"
        "table{border-collapse:collapse;width:100%;font-size:10pt;}"
        "th,td{border:1px solid #b7c2cc;padding:6px 9px;text-align:left;vertical-align:top;}"
        "th{background:#e8eef5;}"
        "a{color:#0b6a9d;text-decoration:none;}"
        "</style></head><body>" + body + "</body></html>"
    )
    try:
        pdf_path = md_path.with_suffix(".pdf")
        HTML(string=html).write_pdf(str(pdf_path))
        log.info("PDF generated: %s", pdf_path)
    except Exception as exc:
        log.warning("PDF generation failed: %s", exc)

# ----------------------------------------------------------------------------
# MAIN PIPELINE
# ----------------------------------------------------------------------------


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Ingredient price intelligence & renegotiation leverage agent")
    parser.add_argument("--ingredients", type=Path, default=DEFAULT_INGREDIENTS_FILE,
                        help="input file (.txt / .md / .json), default: %(default)s")
    parser.add_argument("--out", type=Path, default=REPORTS_DIR,
                        help="output directory for report, PDF and audit JSON")
    parser.add_argument("--backend", choices=LLM_BACKENDS,
                        default=os.getenv("LLM_BACKEND", "ollama"))
    parser.add_argument("--model", default=None, help="override the backend's model name")
    parser.add_argument("--workers", type=int, default=int(os.getenv("MAX_WORKERS", "3")),
                        help="parallel ingredient workers (default 3)")
    parser.add_argument("--no-cache", action="store_true",
                        help="ignore the weekly per-ingredient result cache")
    parser.add_argument("--no-pdf", action="store_true", help="skip PDF export")
    parser.add_argument("-v", "--verbose", action="store_true", help="enable debug logging")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    setup_logging(args.verbose)
    log.info("=== Price Renegotiation Agent ===")

    model = args.model or DEFAULT_MODELS.get(args.backend, "")
    ingredients = load_ingredients(args.ingredients)
    if not ingredients:
        log.warning("No ingredients found in %s. Check your input file.", args.ingredients)
        return
    preview = ", ".join(ingredients[:10]) + (" ..." if len(ingredients) > 10 else "")
    log.info("Loaded %d ingredient(s): %s", len(ingredients), preview)
    log.info("Backend: %s (%s) | search window: %d days | workers: %d",
             args.backend, model, SOURCE_WINDOW_DAYS, args.workers)

    if args.workers > 1:
        with ThreadPoolExecutor(max_workers=args.workers) as pool:
            analyses = list(pool.map(lambda idx: process_ingredient(ingredients[idx], args),
                                     range(len(ingredients))))
    else:
        analyses = [process_ingredient(ingredient, args) for ingredient in ingredients]

    args.out.mkdir(parents=True, exist_ok=True)

    report_path = args.out / "price_negotiation_report.md"
    report_path.write_text(
        generate_report(analyses, {"backend": args.backend, "model": model}),
        encoding="utf-8",
    )
    log.info("Report saved: %s", report_path)

    audit_path = args.out / "analyses.json"
    audit_path.write_text(json.dumps({
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "backend": args.backend,
        "model": model,
        "ingredients_file": str(args.ingredients),
        "analyses": analyses,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    log.info("Audit data saved: %s", audit_path)

    if not args.no_pdf:
        export_pdf(report_path)


if __name__ == "__main__":
    main()
